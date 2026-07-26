from io import BytesIO
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import get_args

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.api.routes.imports import SPECS_BY_SHEET, expected_headers
from app.models.master_data import Park
from app.models.user import User


def headers(client: TestClient) -> dict[str, str]:
    login = client.post("/api/v1/auth/login", json={"username": "demo_admin", "password": "demo-password"})
    access_token = login.json()["data"]["access_token"]
    return {"Authorization": f"Bearer {access_token}"}


def workbook_bytes(park_name: str = "导入园区") -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, spec in SPECS_BY_SHEET.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(expected_headers(spec))
        if sheet_name == "园区档案":
            worksheet.append(["PARK-IMPORT-001", park_name, "导入测试地址", "ACTIVE", "ERP", "PARK-IMPORT-001", "2026-07-23T00:00:00+00:00", "导入测试"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_precheck_confirm_and_duplicate(client: TestClient, demo_user: User, db_session) -> None:
    auth = headers(client)
    files = {"file": ("web-import.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    precheck = client.post("/api/v1/imports/precheck", headers=auth, files=files)

    assert precheck.status_code == 200
    assert precheck.json()["data"]["status"] == "READY_TO_CONFIRM"
    assert precheck.json()["data"]["summary"]["row_count"] == 1
    batch_id = precheck.json()["data"]["batch_id"]

    confirmed = client.post(f"/api/v1/imports/{batch_id}/confirm", headers=auth, json={"confirmed": True})
    assert confirmed.status_code == 200
    assert confirmed.json()["data"]["status"] == "COMPLETED"
    assert db_session.get(Park, "PARK-IMPORT-001").park_name == "导入园区"

    duplicate_precheck = client.post(
        "/api/v1/imports/precheck",
        headers=auth,
        files={"file": ("web-import-duplicate.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    duplicate_batch_id = duplicate_precheck.json()["data"]["batch_id"]
    duplicate_confirmed = client.post(f"/api/v1/imports/{duplicate_batch_id}/confirm", headers=auth, json={"confirmed": True})

    assert duplicate_confirmed.status_code == 200
    results = duplicate_confirmed.json()["data"]["summary"]["results"]
    assert next(item["result"] for item in results if item["sheet_name"] == "园区档案") == "DUPLICATE"


def test_import_precheck_rejects_bad_header(client: TestClient, demo_user: User) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "园区档案"
    worksheet.append(["wrong_header"])
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        "/api/v1/imports/precheck",
        headers=headers(client),
        files={"file": ("bad.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert response.json()["data"]["status"] == "PRECHECK_FAILED"
    assert response.json()["data"]["errors"][0]["error_code"] == "SHEET_MISSING"
    batch_id = response.json()["data"]["batch_id"]
    errors = client.get(f"/api/v1/imports/{batch_id}/errors?page=1&page_size=2", headers=headers(client))
    template = client.get("/api/v1/imports/template", headers=headers(client))

    assert errors.status_code == 200
    assert errors.json()["data"]["total"] > 0
    assert template.status_code == 200
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def full_workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    timestamp = datetime(2026, 7, 23, 0, 0, tzinfo=timezone.utc)
    for sheet_name, spec in SPECS_BY_SHEET.items():
        worksheet = workbook.create_sheet(sheet_name)
        headers = expected_headers(spec)
        worksheet.append(headers)
        values = []
        for field_name, field_info in spec["schema"].model_fields.items():
            if field_name == "remark":
                continue
            if not field_info.is_required() and field_name != "remark":
                values.append(None)
                continue
            choices = [choice for choice in get_args(field_info.annotation) if isinstance(choice, str)]
            if field_name == "remark":
                values.append("导入测试")
            elif field_name == "currency":
                values.append("CNY")
            elif choices:
                values.append(choices[0])
            elif field_name == "line_no" or field_name == "priority" or field_name == "status_version":
                values.append(1)
            elif field_name.endswith("_at") or field_name in {"created_at", "updated_at"}:
                values.append(timestamp.isoformat())
            elif field_name.endswith("_date") or field_name in {"period_start", "period_end", "valid_from", "valid_to"}:
                values.append(date(2026, 7, 23).isoformat())
            elif field_name == "on_duty":
                values.append(True)
            elif field_name == "source_url":
                values.append("https://example.com/source")
            elif field_name == "unit":
                values.append("kg")
            elif field_name.endswith("_id"):
                values.append(f"{field_name.upper()}-IMPORT-001")
            elif field_info.annotation in (int, float, Decimal):
                values.append(1)
            else:
                values.append("导入测试")
        values.extend(["ERP", f"{sheet_name}-IMPORT-001", timestamp.isoformat(), "导入测试"])
        worksheet.append(values)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_full_business_sheet_set(client: TestClient, demo_user: User) -> None:
    response = client.post(
        "/api/v1/imports/precheck",
        headers=headers(client),
        files={"file": ("full-import.xlsx", full_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert response.json()["data"]["status"] == "READY_TO_CONFIRM"
    assert response.json()["data"]["summary"]["row_count"] == len(SPECS_BY_SHEET)
    batch_id = response.json()["data"]["batch_id"]
    confirmed = client.post(f"/api/v1/imports/{batch_id}/confirm", headers=headers(client), json={"confirmed": True})

    assert confirmed.status_code == 200
    results = confirmed.json()["data"]["summary"]["results"]
    assert len(results) == len(SPECS_BY_SHEET)
