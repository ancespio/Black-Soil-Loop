from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Annotated, Any
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pydantic import TypeAdapter, ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.db.session import get_db
from app.models.business_records import Inventory, ReturnRecord, SalesOrderLine
from app.models.imports import ImportBatch
from app.models.master_data import Enterprise, EnterpriseTag, Park, Partner, Store
from app.models.planning_records import Policy, Preorder, ProcurementDemand, SupplierQuote
from app.models.production import Bom, ProductionOrder, ProductionPlan
from app.models.transport import FreezerRecord, TransportResource, TransportTaskSummary
from app.models.operations import ProcurementHistory
from app.models.user import User
from app.schemas.business_records import InventoryCreate, ReturnRecordCreate, SalesOrderLineCreate
from app.schemas.common import ResponseEnvelope, response_envelope
from app.schemas.imports import ImportConfirmRequest
from app.schemas.master_data import EnterpriseCreate, EnterpriseTagCreate, ParkCreate, PartnerCreate, StoreCreate
from app.schemas.planning_records import PolicyCreate, PreorderCreate, ProcurementDemandCreate, SupplierQuoteCreate
from app.schemas.production import BomCreate, ProductionOrderCreate, ProductionPlanCreate
from app.schemas.transport import FreezerRecordCreate, TransportResourceCreate, TransportTaskSummaryCreate
from app.schemas.operations import ProcurementHistoryCreate

router = APIRouter(tags=["E01 Imports"])
DATETIME_ADAPTER = TypeAdapter(datetime)
CONTROL_SHEETS = {"导入说明", "字段字典", "枚举字典", "业务键字典"}

SHEET_SPECS: list[dict[str, Any]] = [
    {"sheet": "园区档案", "model": Park, "schema": ParkCreate, "keys": ("park_id",)},
    {"sheet": "企业档案", "model": Enterprise, "schema": EnterpriseCreate, "keys": ("enterprise_id",)},
    {"sheet": "企业标签", "model": EnterpriseTag, "schema": EnterpriseTagCreate, "keys": ("enterprise_id", "tag")},
    {"sheet": "合作方", "model": Partner, "schema": PartnerCreate, "keys": ("partner_id",)},
    {"sheet": "门店", "model": Store, "schema": StoreCreate, "keys": ("store_id",)},
    {"sheet": "生产计划", "model": ProductionPlan, "schema": ProductionPlanCreate, "keys": ("plan_id",)},
    {"sheet": "生产订单", "model": ProductionOrder, "schema": ProductionOrderCreate, "keys": ("production_order_id",)},
    {"sheet": "物料清单", "model": Bom, "schema": BomCreate, "keys": ("bom_id", "product_id", "material_id")},
    {"sheet": "企业库存", "model": Inventory, "schema": InventoryCreate, "keys": ("inventory_record_id",)},
    {"sheet": "销售订单明细", "model": SalesOrderLine, "schema": SalesOrderLineCreate, "keys": ("sales_order_id", "line_no")},
    {"sheet": "退货记录", "model": ReturnRecord, "schema": ReturnRecordCreate, "keys": ("return_id",)},
    {"sheet": "运输任务摘要", "model": TransportTaskSummary, "schema": TransportTaskSummaryCreate, "keys": ("task_id",)},
    {"sheet": "运输资源", "model": TransportResource, "schema": TransportResourceCreate, "keys": ("driver_id", "vehicle_id")},
    {"sheet": "冻库记录", "model": FreezerRecord, "schema": FreezerRecordCreate, "keys": ("freezer_id", "recorded_at")},
    {"sheet": "预订单", "model": Preorder, "schema": PreorderCreate, "keys": ("preorder_id",)},
    {"sheet": "采购需求", "model": ProcurementDemand, "schema": ProcurementDemandCreate, "keys": ("demand_id",)},
    {"sheet": "供应商报价", "model": SupplierQuote, "schema": SupplierQuoteCreate, "keys": ("supplier_id", "material_id", "tier_id")},
    {"sheet": "采购历史", "model": ProcurementHistory, "schema": ProcurementHistoryCreate, "keys": ("purchase_record_id",)},
    {"sheet": "政策记录", "model": Policy, "schema": PolicyCreate, "keys": ("policy_id",)},
]
SPECS_BY_SHEET = {spec["sheet"]: spec for spec in SHEET_SPECS}


class ImportConflict(Exception):
    def __init__(self, errors: list[dict[str, Any]]):
        super().__init__("导入存在冲突")
        self.errors = errors


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def expected_headers(spec: dict[str, Any]) -> list[str]:
    fields = [field for field in spec["schema"].model_fields if field != "remark"]
    return fields + ["source_system", "source_record_id", "source_updated_at", "remark"]


def cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def error_item(sheet_name: str, row_number: int | None, field_key: str | None, error_code: str, message: str) -> dict[str, Any]:
    return {
        "sheet_name": sheet_name,
        "row_number": row_number,
        "field_key": field_key,
        "error_code": error_code,
        "message": message,
    }


def canonical(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def payload_matches(record: Any, schema: type, payload: dict[str, Any]) -> bool:
    for field in schema.model_fields:
        if canonical(getattr(record, field)) != canonical(payload.get(field)):
            return False
    return True


def check_permission(user: User, schema: type, payload: dict[str, Any], sheet_name: str, row_number: int) -> dict[str, Any] | None:
    if user.role != "enterprise_admin":
        return None
    if "enterprise_id" not in schema.model_fields or payload.get("enterprise_id") is None:
        return error_item(sheet_name, row_number, "enterprise_id", "FORBIDDEN", "企业管理员只能导入带自身 enterprise_id 的记录")
    if payload["enterprise_id"] not in (user.enterprise_ids or []):
        return error_item(sheet_name, row_number, "enterprise_id", "FORBIDDEN", "超出当前企业授权范围")
    return None


def validate_row(spec: dict[str, Any], raw: dict[str, Any], sheet_name: str, row_number: int, user: User) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    source_system = raw.get("source_system")
    source_record_id = raw.get("source_record_id")
    source_updated_at = raw.get("source_updated_at")
    if source_system is None:
        errors.append(error_item(sheet_name, row_number, "source_system", "REQUIRED", "来源系统不能为空"))
    if source_record_id is None:
        errors.append(error_item(sheet_name, row_number, "source_record_id", "REQUIRED", "来源记录编号不能为空"))
    try:
        source_updated = DATETIME_ADAPTER.validate_python(source_updated_at)
    except ValidationError as exc:
        errors.append(error_item(sheet_name, row_number, "source_updated_at", "INVALID_DATETIME", str(exc.errors()[0]["msg"])))
        source_updated = None
    payload_values = {field: raw.get(field) for field in spec["schema"].model_fields}
    try:
        payload = spec["schema"].model_validate(payload_values)
    except ValidationError as exc:
        for validation_error in exc.errors():
            location = validation_error.get("loc", ("",))
            field_key = str(location[0]) if location else None
            errors.append(error_item(sheet_name, row_number, field_key, "VALIDATION_ERROR", validation_error["msg"]))
        payload = None
    if payload is not None:
        permission_error = check_permission(user, spec["schema"], payload.model_dump(), sheet_name, row_number)
        if permission_error is not None:
            errors.append(permission_error)
    if errors or payload is None or source_updated is None or source_system is None or source_record_id is None:
        return None, errors
    return {
        "sheet_name": sheet_name,
        "row_number": row_number,
        "payload": payload.model_dump(mode="json"),
        "source_system": str(source_system),
        "source_record_id": str(source_record_id),
        "source_updated_at": source_updated.isoformat(),
    }, errors


def read_workbook(content: bytes, user: User) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    errors: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    actual_sheets = set(workbook.sheetnames)
    expected_sheets = set(SPECS_BY_SHEET)
    for missing in sorted(expected_sheets - actual_sheets):
        errors.append(error_item(missing, None, None, "SHEET_MISSING", "缺少业务工作表"))
    for unknown in sorted(actual_sheets - expected_sheets - CONTROL_SHEETS):
        errors.append(error_item(unknown, None, None, "SHEET_UNKNOWN", "存在未被模板契约允许的工作表"))
    for sheet_name, spec in SPECS_BY_SHEET.items():
        if sheet_name not in actual_sheets:
            continue
        worksheet = workbook[sheet_name]
        values = worksheet.iter_rows(values_only=True)
        header = [cell_value(value) for value in next(values, ())]
        while header and header[-1] is None:
            header.pop()
        expected = expected_headers(spec)
        if header != expected:
            errors.append(error_item(sheet_name, 1, None, "HEADER_MISMATCH", f"表头必须严格匹配模板，期望: {','.join(expected)}"))
            continue
        for row_number, raw_values in enumerate(values, start=2):
            values_list = list(raw_values)
            if not any(cell_value(value) is not None for value in values_list):
                continue
            raw = {field: cell_value(values_list[index] if index < len(values_list) else None) for index, field in enumerate(header)}
            row, row_errors = validate_row(spec, raw, sheet_name, row_number, user)
            errors.extend(row_errors)
            if row is not None:
                rows.append(row)
    workbook.close()
    summary = {
        "sheet_count": len(SHEET_SPECS),
        "row_count": len(rows),
        "error_count": len(errors),
        "sheets": {spec["sheet"]: sum(1 for row in rows if row["sheet_name"] == spec["sheet"]) for spec in SHEET_SPECS},
    }
    return rows, errors, summary


def batch_data(batch: ImportBatch) -> dict[str, Any]:
    return {
        "batch_id": batch.batch_id,
        "file_name": batch.file_name,
        "uploaded_by": batch.uploaded_by,
        "status": batch.status,
        "uploaded_at": batch.uploaded_at,
        "confirmed_at": batch.confirmed_at,
        "summary": batch.summary_json,
        "errors": batch.errors_json,
    }


def find_by_business_key(db: Session, spec: dict[str, Any], payload: dict[str, Any]) -> Any | None:
    conditions = [getattr(spec["model"], key) == payload[key] for key in spec["keys"]]
    return db.scalar(select(spec["model"]).where(*conditions))


def apply_import_row(db: Session, spec: dict[str, Any], row: dict[str, Any]) -> str:
    schema = spec["schema"]
    payload_model = schema.model_validate(row["payload"])
    payload = payload_model.model_dump()
    source_updated_at = DATETIME_ADAPTER.validate_python(row["source_updated_at"])
    model = spec["model"]
    source_record = db.scalar(select(model).where(model.source_system == row["source_system"], model.source_record_id == row["source_record_id"]))
    business_record = find_by_business_key(db, spec, payload)
    if source_record is not None and business_record is not None and source_record is not business_record:
        raise ImportConflict([error_item(row["sheet_name"], row["row_number"], None, "IDEMPOTENCY_CONFLICT", "来源记录和业务键分别指向不同实体")])
    record = source_record or business_record
    if record is not None:
        existing_time = record.source_updated_at
        if existing_time.tzinfo is None:
            existing_time = existing_time.replace(tzinfo=timezone.utc)
        if source_updated_at < existing_time:
            return "SKIPPED_STALE"
        same_content = payload_matches(record, schema, payload)
        if source_updated_at == existing_time:
            if same_content:
                return "DUPLICATE"
            raise ImportConflict([error_item(row["sheet_name"], row["row_number"], None, "IDEMPOTENCY_CONFLICT", "同一实体同一 source_updated_at 的内容不同")])
        for field, value in payload.items():
            setattr(record, field, value)
        record.source_system = row["source_system"]
        record.source_record_id = row["source_record_id"]
        record.source_updated_at = source_updated_at
        record.object_version += 1
        return "UPDATED"
    record = model(
        **payload,
        source_system=row["source_system"],
        source_record_id=row["source_record_id"],
        source_updated_at=source_updated_at,
    )
    db.add(record)
    return "CREATED"


def ensure_batch_access(batch: ImportBatch, user: User) -> None:
    if user.role != "park_admin" and batch.uploaded_by != user.user_id:
        raise HTTPException(status_code=403, detail={"code": "FORBIDDEN", "message": "无权访问该导入批次"})


@router.get("/imports/template", response_class=FileResponse)
def download_import_template(
    user: Annotated[User, Depends(get_current_user)],
) -> FileResponse:
    del user
    template = Path(__file__).resolve().parents[3] / "web-import-template-v0.1.xlsx"
    if not template.exists():
        raise HTTPException(status_code=404, detail={"code": "NOT_FOUND", "message": "导入模板不存在"})
    return FileResponse(
        template,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="web-import-template-v0.1.xlsx",
    )


@router.post("/imports/precheck", response_model=ResponseEnvelope[dict])
def precheck_import(
    request: Request,
    file: Annotated[UploadFile, File()],
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> dict:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail={"code": "VALIDATION_ERROR", "message": "仅支持 .xlsx 文件"})
    content = file.file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail={"code": "VALIDATION_ERROR", "message": "工作簿不能超过 20 MB"})
    try:
        rows, errors, summary = read_workbook(content, user)
    except Exception as exc:
        raise HTTPException(status_code=400, detail={"code": "VALIDATION_ERROR", "message": f"工作簿无法读取: {exc}"}) from exc
    batch_id = f"IMPORT-{uuid4().hex[:16].upper()}"
    batch = ImportBatch(
        batch_id=batch_id,
        file_name=file.filename,
        uploaded_by=user.user_id,
        status="READY_TO_CONFIRM" if not errors else "PRECHECK_FAILED",
        uploaded_at=now_utc(),
        rows_json=rows,
        summary_json=summary,
        errors_json=errors,
    )
    db.add(batch)
    db.commit()
    code = "READY_TO_CONFIRM" if not errors else "PRECHECK_FAILED"
    status = "PROCESSED" if not errors else "REJECTED"
    return response_envelope(batch_data(batch), status=status, code=code, trace_id=request.state.trace_id)


@router.get("/imports/{batch_id}", response_model=ResponseEnvelope[dict])
def get_import_batch(
    request: Request,
    batch_id: str,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> dict:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail={"code": "NOT_FOUND", "message": "导入批次不存在"})
    ensure_batch_access(batch, user)
    return response_envelope(batch_data(batch), trace_id=request.state.trace_id)


@router.get("/imports/{batch_id}/errors", response_model=ResponseEnvelope[dict])
def list_import_errors(
    request: Request,
    batch_id: str,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
) -> dict:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail={"code": "NOT_FOUND", "message": "导入批次不存在"})
    ensure_batch_access(batch, user)
    errors = batch.errors_json or []
    start = (page - 1) * page_size
    data = {"items": errors[start : start + page_size], "total": len(errors), "page": page, "page_size": page_size}
    return response_envelope(data, trace_id=request.state.trace_id)


@router.post("/imports/{batch_id}/confirm", response_model=ResponseEnvelope[dict])
def confirm_import(
    request: Request,
    batch_id: str,
    body: ImportConfirmRequest,
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(get_current_user)],
) -> dict:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail={"code": "NOT_FOUND", "message": "导入批次不存在"})
    ensure_batch_access(batch, user)
    if batch.status != "READY_TO_CONFIRM":
        raise HTTPException(status_code=409, detail={"code": "IMPORT_STATE_CONFLICT", "message": f"当前批次状态为 {batch.status}，不可确认"})
    batch.status = "IMPORTING"
    db.commit()
    results: list[dict[str, Any]] = []
    try:
        for spec in SHEET_SPECS:
            for row in [row for row in batch.rows_json if row["sheet_name"] == spec["sheet"]]:
                result = apply_import_row(db, spec, row)
                results.append({"sheet_name": row["sheet_name"], "row_number": row["row_number"], "result": result})
        db.flush()
        batch.status = "COMPLETED"
        batch.confirmed_at = now_utc()
        batch.summary_json = {**batch.summary_json, "results": results}
        db.commit()
    except ImportConflict as exc:
        db.rollback()
        batch = db.get(ImportBatch, batch_id)
        batch.status = "FAILED"
        batch.errors_json = exc.errors
        batch.error_message = "导入存在冲突，整本未写入"
        db.commit()
        raise HTTPException(status_code=409, detail={"code": "IDEMPOTENCY_CONFLICT", "message": "导入存在冲突，整本未写入"}) from exc
    except Exception as exc:
        db.rollback()
        batch = db.get(ImportBatch, batch_id)
        batch.status = "FAILED"
        batch.errors_json = [error_item("", None, None, "IMPORT_FAILED", str(exc))]
        batch.error_message = "导入失败，整本未写入"
        db.commit()
        raise HTTPException(status_code=400, detail={"code": "IMPORT_FAILED", "message": "导入失败，整本未写入"}) from exc
    return response_envelope(batch_data(batch), trace_id=request.state.trace_id)
